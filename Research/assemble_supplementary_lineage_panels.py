"""Assemble the unchanged U1.2 and U1.3 CNV panels with separate captions."""

from pathlib import Path
import argparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt


SOURCE = OUTPUT = None

CAPTIONS = {
    "A": (
        "U1.2 lineage. The chr1–chr22 matrix reports the summed length (Mb) of "
        "filtered differential CNV calls for U1→U1.2 (orange), U1.2→U1.2.1 "
        "(blue) and U1.2→U1.2.2 (green). Darker cells indicate greater summed "
        "length; unfilled cells indicate no calls. Outlined chr1, chr3, chr7 "
        "and chr15 are enlarged below; colored bars mark the genomic positions "
        "of the original pairwise calls (Mb). Colors identify comparisons, "
        "not gain or loss."
    ),
    "B": (
        "U1.3 lineage. The chr1–chr22 matrix reports the summed length (Mb) of "
        "filtered differential CNV calls for U1→U1.3 (orange), U1.3→U1.3.1 "
        "(blue) and U1.3→U1.3.2 (green). Darker cells indicate greater summed "
        "length; unfilled cells indicate no calls. Outlined chr1, chr2, chr3, "
        "chr7, chr9, chr17 and chr21 correspond to the numbered positional "
        "tracks below. Colored bars mark the genomic positions of the original "
        "pairwise calls (Mb); colors identify comparisons, not gain or loss."
    ),
}


def set_font(run, size, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold


def main():
    global SOURCE, OUTPUT
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(); SOURCE = args.input_dir; OUTPUT = args.output
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(185)
    section.page_height = Mm(475)
    section.left_margin = Mm(5)
    section.right_margin = Mm(5)
    section.top_margin = Mm(5)
    section.bottom_margin = Mm(5)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for label, filename in (("A", "CNV_inheritance_U12.png"),
                            ("B", "CNV_inheritance_U13.png")):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0 if label == "A" else 5)
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), 9, bold=True)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SOURCE / filename), width=Mm(175))

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(f"{label}. "), 8, bold=True)
        set_font(p.add_run(CAPTIONS[label]), 8)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
